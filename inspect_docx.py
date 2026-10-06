
from docx import Document
import sys

d = Document(r"d:\Programming\Arduino Stuff\rpi5 potato sorter\Micro Final Project Report.docx")

out = []
out.append("===")
out.append(f"Total paragraphs: {len(d.paragraphs)}")
out.append(f"Total tables: {len(d.tables)}")
out.append(f"Total sections: {len(d.sections)}")
out.append("===")

for i, p in enumerate(d.paragraphs):
    text = p.text
    if not text and not p.runs:
        marker = "[EMPTY PARAGRAPH]"
    else:
        marker = repr(text[:300])
    style = p.style.name if p.style else "None"
    align = str(p.alignment) if p.alignment else "None"
    out.append(f"[{i:03d}] style={style:<20s} align={align:<10s} text={marker}")

out.append("\n\n=== TABLES ===")
for ti, t in enumerate(d.tables):
    out.append(f"\nTable {ti}: {len(t.rows)} rows x {len(t.columns)} cols")
    for ri, row in enumerate(t.rows):
        cells = [c.text.strip() for c in row.cells]
        out.append(f"  Row {ri}: {cells}")

out.append("\n\n=== RUN FONT DETAILS FOR FIRST 40 PARAGRAPHS (only non-empty) ===")
for i, p in enumerate(d.paragraphs[:40]):
    if not p.text.strip():
        continue
    run_info = []
    for r in p.runs[:3]:
        name = r.font.name
        size = r.font.size.pt if r.font.size else None
        bold = r.bold
        run_info.append(f"font={name} size={size} bold={bold}")
    out.append(f"[{i:03d}] {repr(p.text[:100])} => runs: {run_info}")

with open(r"d:\Programming\Arduino Stuff\rpi5 potato sorter\docx_inspect_output.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(out))

print("WROTE", len(out), "lines")
